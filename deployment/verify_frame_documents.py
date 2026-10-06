"""Independent verification of actual Rust PDF/Excel outputs, including page bounds."""
import json
import re
import subprocess
import sys
from pathlib import Path

import pdfplumber
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
folder = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "qa/readiness/frame-documents"
checks = []


def check(condition, description):
    if not condition:
        raise AssertionError(description)
    checks.append(description)


def document(name):
    reader = PdfReader(folder / name)
    check(len(reader.pages) > 0, f"{name}: readable nonempty PDF")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    with pdfplumber.open(folder / name) as pdf:
        for index, page in enumerate(pdf.pages):
            for char in page.chars:
                # Font ascent may extend above the text cursor. Allow a one
                # point numerical tolerance, not text outside the page.
                assert -1 <= char["x0"] <= char["x1"] <= page.width + 1, (name, index, char)
                assert -1 <= char["top"] <= char["bottom"] <= page.height + 1, (name, index, char)
    check(True, f"{name}: every character stays inside its page")
    return reader, text


schedule, text = document("schedule-long.pdf")
check(len(schedule.pages) >= 3, "Schedule paginates the long 78-frame fixture")
for number in range(1, 79):
    assert len(re.findall(rf"\bQ{number:03}\b", text)) == 1, f"Missing/duplicate schedule frame Q{number:03}"
check(True, "Schedule retains all 78 unique frame marks")
check(all(value in text for value in ["Café", "façade", "Ελληνικά", "Кириллица", "Tilt and turn", "HR++"]),
      "Schedule preserves Unicode, panel types and glazing information")

production, text = document("production-long.pdf")
for number in range(1, 91):
    assert len(re.findall(rf"\bCUT-{number:03}\b", text)) == 1, f"Missing/duplicate cut row CUT-{number:03}"
check(len(production.pages) > 1, "Production list paginates and retains all 90 cutting rows")
check(all(value in text for value in ["Glass list", "Hardware list", "Gasket list", "Glazing bead list", "Bill of materials"]),
      "Production PDF includes every populated production section")

estimate, text = document("estimate-long.pdf")
for number in range(1, 79):
    assert len(re.findall(rf"\bQ{number:03}\b", text)) == 1, f"Missing estimate item Q{number:03}"
for number in range(1, 71):
    assert len(re.findall(rf"\bTERM-{number:03}\b", text)) == 1, f"Missing estimate term TERM-{number:03}"
check(len(estimate.pages) > 1, "Estimate retains 78 items and all 70 terms across pages")
check("9227.01" in text, "Estimate total matches 78 x EUR 100.25 with 18% tax")

labels, text = document("labels.pdf")
for number in range(1, 9):
    for suffix in ["SL", "SR", "DB", "DO", "V1-RSL", "V1-RSR", "V1-RDB", "V1-RDO"]:
        assert f"Q{number:03}-{suffix}" in text, f"Missing actual piece label Q{number:03}-{suffix}"
check(len(labels.pages) >= 3, "Labels retain the actual cutting/CNC piece IDs on multiple sheets")
workshop, text = document("workshop.pdf")
check(len(workshop.pages) == 1 and "1200.25 mm" in text and "1500.5 mm" in text,
      "Workshop A3 sheet displays the actual fractional dimensions")
wide, text = document("workshop-wide.pdf")
check(len(wide.pages) == 1 and "120000.25 mm" in text and "10000.5 mm" in text,
      "Wide workshop sheet retains both dimension labels inside its page")

book = load_workbook(folder / "schedule-long.xlsx", data_only=False)
sheet = book["Frame schedule"]
check(sheet.max_row == 79 and sheet["A79"].value == "Q078", "Excel schedule retains all 78 frame rows")
check(sheet["C2"].value == 1200.25 and sheet["D2"].value == 1500.5, "Excel schedule preserves fractional dimensions")
book = load_workbook(folder / "production-long.xlsx", data_only=False)
check(book["Cut list"].max_row == 91 and book["Cut list"]["B91"].value == "CUT-090",
      "Excel production workbook retains all 90 cutting rows")
check(book["Cut list"]["F2"].value % 1 != 0, "Excel cut lengths preserve fractional precision")
check({"Cut list", "Glass list", "Hardware list", "Gasket list", "Glazing bead list", "Bill of materials"} <= set(book.sheetnames),
      "Excel workbook includes every populated production section")
panel, text = document("panel-production.pdf")
check("Panel list" in text and "Sandwich panel" in text and "PANEL-01-V1" in text,
      "Production PDF includes the actual sandwich panel position and material")
book = load_workbook(folder / "panel-production.xlsx", data_only=False)
check(book["Panel list"].max_row == 2 and book["Panel list"]["E2"].value == "Sandwich panel",
      "Production workbook retains the actual sandwich panel row")

render = folder / "rendered"
render.mkdir(exist_ok=True)
names = ["schedule-long", "production-long", "estimate-long", "labels", "workshop", "workshop-wide", "panel-production"]
if "--browser" in sys.argv:
    schedule, text = document("browser-schedule.pdf")
    for number in range(1, 79):
        assert len(re.findall(rf"\bQ{number:03}\b", text)) == 1, f"Missing browser schedule frame Q{number:03}"
    check(len(schedule.pages) >= 3, "Actual browser schedule retains all 78 frames across pages")
    book = load_workbook(folder / "browser-schedule.xlsx", data_only=False)
    check(book["Frame schedule"].max_row == 79 and book["Frame schedule"]["C2"].value == 1200.25,
          "Actual browser Excel schedule retains all frames and fractional dimensions")
    production, text = document("browser-production.pdf")
    for number in range(1, 9):
        for suffix in ["SL", "SR", "DB", "DO", "V1-RSL", "V1-RSR", "V1-RDB", "V1-RDO"]:
            assert len(re.findall(rf"\bQ{number:03}-{suffix}\b", text)) == 1, f"Missing/duplicate browser cut Q{number:03}-{suffix}"
    check(all(value in text for value in ["Hardware list", "Gasket list", "Glazing bead list", "Bill of materials"]),
          "Actual browser production PDF retains all frame cuts and populated sections")
    book = load_workbook(folder / "browser-production.xlsx", data_only=False)
    check(book["Cut list"].max_row == 65 and book["Glass list"].max_row == 9,
          "Actual browser production workbook retains all 64 pieces and eight glazing rows")
    labels, text = document("browser-labels.pdf")
    for number in range(1, 9):
        for suffix in ["SL", "SR", "DB", "DO", "V1-RSL", "V1-RSR", "V1-RDB", "V1-RDO"]:
            assert len(re.findall(rf"\bQ{number:03}-{suffix}\b", text)) == 1, f"Missing/duplicate browser label Q{number:03}-{suffix}"
    check(len(labels.pages) == 4, "Actual browser label sheets retain all 64 actual piece identifiers")
    workshop, text = document("browser-workshop.pdf")
    check(len(workshop.pages) == 1 and "1200.25 mm" in text and "1500.5 mm" in text,
          "Actual browser workshop sheet retains its fractional frame dimensions")
    estimate, text = document("browser-estimate.pdf")
    for number in range(1, 9):
        assert len(re.findall(rf"\bQ{number:03}\b", text)) == 1, f"Missing browser estimate item Q{number:03}"
    evidence = json.loads((root / "qa/readiness/frame-document-downloads.json").read_text())
    check(f'{evidence["referenceEstimateTotal"]:.2f}' in text,
          "Actual browser estimate total agrees with the saved reference draft")
    names += ["browser-schedule", "browser-production", "browser-labels", "browser-workshop", "browser-estimate"]

for name in names:
    subprocess.run(["pdftoppm", "-r", "90", "-png", str(folder / f"{name}.pdf"), str(render / name)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)

result = {"status": "passed", "checks": checks, "renderedPages": len(list(render.glob("*.png")))}
(root / "qa/readiness/frame-document-validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
